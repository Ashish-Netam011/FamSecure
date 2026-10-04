# -*- coding: utf-8 -*-
"""
Builds the official FarmRescue presentation deck.

Usage:  python presentation/build_deck.py
Output: presentation/FarmRescue-Official-Deck.pptx

Requires: python-pptx, Pillow (both already available in this environment).
No network access, no other dependencies.
"""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, "shots")
OUT = os.path.join(HERE, "FarmRescue-Official-Deck.pptx")

# ----------------------------------------------------------------------------
# Theme
# ----------------------------------------------------------------------------
FOREST = RGBColor(0x0B, 0x2E, 0x1B)   # deepest green (title bg)
DARK = RGBColor(0x13, 0x3B, 0x27)     # header band
GREEN = RGBColor(0x15, 0x80, 0x3D)    # primary green-700
EMERALD = RGBColor(0x10, 0xB9, 0x81)  # accent
MINT = RGBColor(0xEC, 0xFD, 0xF5)     # panel tint
LIGHT = RGBColor(0xF1, 0xF5, 0xF9)    # light slate
BORDER = RGBColor(0xCB, 0xD5, 0xE1)
SLATE = RGBColor(0x1E, 0x29, 0x3B)
GRAY = RGBColor(0x5B, 0x6B, 0x80)
RED = RGBColor(0xDC, 0x26, 0x26)
AMBER = RGBColor(0xD9, 0x77, 0x06)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xD1, 0xFA, 0xE5)

HEAD = "Segoe UI Semibold"
BODY = "Segoe UI"

SW, SH = 13.333, 7.5
M = 0.62


# ----------------------------------------------------------------------------
# Low-level helpers
# ----------------------------------------------------------------------------
def slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg(s, color):
    f = s.background.fill
    f.solid()
    f.fore_color.rgb = color


def bg_gradient(s, c1, c2, angle=45):
    f = s.background.fill
    f.gradient()
    try:
        f.gradient_angle = angle
        stops = f.gradient_stops
        stops[0].color.rgb = c1
        stops[0].position = 0.0
        stops[1].color.rgb = c2
        stops[1].position = 1.0
    except Exception:
        f.solid()
        f.fore_color.rgb = c1


def shape(s, x, y, w, h, fill=None, line=None, line_w=1.0,
          kind=MSO_SHAPE.RECTANGLE, radius=None):
    sh = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    sh.shadow.inherit = False
    if radius is not None:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    sh.text_frame.word_wrap = True
    return sh


def card(s, x, y, w, h, fill=WHITE, line=BORDER, radius=0.06):
    return shape(s, x, y, w, h, fill=fill, line=line, line_w=0.75,
                 kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=radius)


def text(s, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def par(tf, txt, size=12, color=SLATE, bold=False, font=BODY,
        align=PP_ALIGN.LEFT, after=4, before=0, first=False, line=None,
        italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.space_after = Pt(after)
    p.space_before = Pt(before)
    if line:
        p.line_spacing = line
    r = p.add_run()
    r.text = txt
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    r.font.name = font
    return p


def bullets(tf, items, size=11.5, color=SLATE, after=7, first=True, gap=0.16):
    """items: list of str or (str, bool_bold)"""
    started = not first
    for it in items:
        bold = False
        txt = it
        if isinstance(it, tuple):
            txt, bold = it
        par(tf, "\u2022  " + txt, size=size, color=color, bold=bold,
            after=after, first=not started, line=1.05)
        started = True


def img_fit(s, path, x, y, w, h, border=True):
    with Image.open(path) as im:
        iw, ih = im.size
    ar = iw / float(ih)
    if (w / h) > ar:
        nh = h
        nw = h * ar
    else:
        nw = w
        nh = w / ar
    nx = x + (w - nw) / 2.0
    ny = y + (h - nh) / 2.0
    if border:
        shape(s, nx - 0.035, ny - 0.035, nw + 0.07, nh + 0.07,
              fill=WHITE, line=BORDER, line_w=0.75,
              kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.02)
    s.shapes.add_picture(path, Inches(nx), Inches(ny), Inches(nw), Inches(nh))


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


SLIDE_NO = {"n": 0}


def header(s, kicker, title, total=None):
    shape(s, 0, 0, SW, 1.18, fill=DARK)
    shape(s, 0, 1.18, SW, 0.055, fill=EMERALD)
    tf = text(s, M, 0.16, SW - 2 * M - 1.4, 0.34)
    par(tf, kicker.upper(), size=10.5, color=EMERALD, bold=True, first=True,
        after=0)
    tf2 = text(s, M, 0.46, SW - 2 * M - 1.4, 0.62)
    par(tf2, title, size=25, color=WHITE, bold=True, font=HEAD, first=True,
        after=0)
    # page number
    SLIDE_NO["n"] += 1
    tfp = text(s, SW - M - 1.2, 0.34, 1.2, 0.4)
    par(tfp, "%02d" % SLIDE_NO["n"], size=20, color=RGBColor(0x6E, 0xE7, 0xB7),
        bold=True, font=HEAD, align=PP_ALIGN.RIGHT, first=True, after=0)


def footer(s, txt="FarmRescue  \u00b7  Emergency Agricultural Early-Warning & Response Platform"):
    tf = text(s, M, SH - 0.42, SW - 2 * M, 0.26)
    par(tf, txt, size=8.5, color=GRAY, first=True, after=0)


def chip(s, x, y, w, h, num, label, fill=MINT, accent=GREEN):
    card(s, x, y, w, h, fill=fill, line=None, radius=0.12)
    tf = text(s, x + 0.18, y + 0.14, w - 0.36, 0.5)
    par(tf, num, size=24, color=accent, bold=True, font=HEAD, first=True,
        after=0)
    tf2 = text(s, x + 0.18, y + 0.66, w - 0.36, h - 0.72)
    par(tf2, label, size=10, color=SLATE, first=True, after=0, line=1.0)


def stat_strip(s, y, items):
    """items: list of (num, label)"""
    n = len(items)
    gap = 0.22
    w = (SW - 2 * M - gap * (n - 1)) / n
    for i, (num, label) in enumerate(items):
        chip(s, M + i * (w + gap), y, w, 1.32, num, label)


# ----------------------------------------------------------------------------
# Slide builders
# ----------------------------------------------------------------------------
def slide_title(prs):
    s = slide(prs)
    bg_gradient(s, FOREST, RGBColor(0x05, 0x1A, 0x10), 45)
    # decorative shapes
    shape(s, 10.6, -1.5, 4.6, 4.6, fill=RGBColor(0x12, 0x4A, 0x2E),
          kind=MSO_SHAPE.OVAL)
    shape(s, -1.4, 4.4, 3.8, 3.8, fill=RGBColor(0x0F, 0x3D, 0x26),
          kind=MSO_SHAPE.OVAL)
    shape(s, 3.9, 2.02, 0.9, 0.075, fill=EMERALD)

    tf = text(s, M + 0.2, 0.62, 10.5, 0.4)
    par(tf, "OFFICIAL PROJECT BRIEFING  \u00b7  OCTOBER 2026", size=12,
        color=EMERALD, bold=True, first=True, after=0)

    tf = text(s, M + 0.2, 1.05, 11.0, 1.3)
    par(tf, "\U0001F33E  FarmRescue", size=54, color=WHITE, bold=True,
        font=HEAD, first=True, after=0)

    tf = text(s, M + 0.2, 2.30, 11.4, 0.85)
    par(tf, "Emergency Agricultural Early-Warning & Response Platform",
        size=23, color=PALE, font=HEAD, first=True, after=4)
    par(tf, "\u201cWhat should I do when the farm suddenly goes wrong?\u201d",
        size=15, color=RGBColor(0x9C, 0xD8, 0xB8), italic=True, after=0)

    tf = text(s, M + 0.2, 3.95, 11.4, 0.9)
    par(tf, "Deterministic agronomic risk intelligence that turns raw field "
            "evidence into validated, expert-authored action plans in seconds \u2014",
        size=14, color=WHITE, first=True, after=3)
    par(tf, "where data and rules decide, and AI only explains.",
        size=14, color=WHITE, after=0)

    # bottom meta bar
    shape(s, 0, SH - 1.05, SW, 1.05, fill=RGBColor(0x04, 0x14, 0x0C))
    tf = text(s, M + 0.2, SH - 0.80, 8.6, 0.6)
    par(tf, "Next.js 14 (App Router)  \u00b7  TypeScript  \u00b7  Tailwind CSS  "
            "\u00b7  MongoDB  \u00b7  Gemini explanation layer",
        size=11.5, color=PALE, first=True, after=0)
    tf = text(s, SW - 4.6, SH - 0.80, 4.0, 0.6)
    par(tf, "Hackathon Product Demo  \u00b7  Prototype Grade: Demo-Ready",
        size=11.5, color=EMERALD, bold=True, align=PP_ALIGN.RIGHT, first=True,
        after=0)
    notes(s, "Opening. One-line pitch: FarmRescue answers the 3 a.m. question a "
             "farmer asks when the crop suddenly goes wrong. Emphasise that the "
             "system is deterministic by design and that AI never decides.")
    return s


def slide_summary(prs):
    s = slide(prs)
    bg(s, WHITE)
    header(s, "Overview", "Executive Summary")

    rows = [
        ("The problem", "Crop-loss emergencies unfold in hours, but trusted, "
                        "location-specific agronomic guidance reaches farmers in days.",
         RED),
        ("The solution", "A field-evidence early-warning platform that assesses a "
                         "reported problem in seconds and returns a specific, "
                         "safety-checked response plan.", GREEN),
        ("The difference", "Data and validated rules decide risk and allowed "
                           "actions. AI only explains validated results \u2014 it never "
                           "overrides a classification or invents a plan.", AMBER),
    ]
    w = SW - 2 * M
    for i, (t, d, accent) in enumerate(rows):
        y = 1.55 + i * 1.32
        card(s, M, y, w, 1.14, fill=LIGHT, line=None)
        shape(s, M, y + 0.14, 0.075, 0.86, fill=accent)
        tf = text(s, M + 0.32, y + 0.16, 3.0, 0.5)
        par(tf, t, size=14, color=SLATE, bold=True, font=HEAD, first=True, after=0)
        tf = text(s, M + 0.32, y + 0.52, w - 0.64, 0.55)
        par(tf, d, size=11.5, color=GRAY, first=True, after=0, line=1.06)

    tf = text(s, M, 5.72, w, 0.35)
    par(tf, "Implementation at a glance", size=12, color=SLATE, bold=True,
        font=HEAD, first=True, after=0)
    stat_strip(s, 6.10, [
        ("31", "automated tests passing (zero-dependency Node runner)"),
        ("12", "expert-authored action plans (ICAR / IRRI-aligned)"),
        ("19", "documented REST API routes"),
        ("0", "decisions authored by the AI layer"),
    ])
    footer(s)
    notes(s, "Three-beat summary. The '0 decisions authored by AI' stat is the "
             "headline governance claim \u2014 be ready to defend it.")
    return s


def slide_problem(prs):
    s = slide(prs)
    bg(s, WHITE)
    header(s, "Context", "The Problem")

    tf = text(s, M, 1.48, SW - 2 * M, 0.8)
    par(tf, "A crop emergency does not wait for the extension officer.",
        size=22, color=DARK, bold=True, font=HEAD, first=True, after=4)
    par(tf, "Waterlogging, heat stress and outbreaks escalate within a single "
            "day \u2014 the window to act is measured in hours.",
        size=12.5, color=GRAY, after=0)

    cards = [
        ("Invisible onset", "Waterlogging, submergence and heat stress cross "
         "critical thresholds silently overnight. By the time damage is visible, "
         "the highest-value intervention window has already closed.", RED),
        ("Generic, sometimes dangerous advice", "Blanket advisory (apply urea, "
         "spray on schedule) is unsafe under active flooding. Advice that ignores "
         "field state can make the loss worse \u2014 e.g. granular urea in standing "
         "water.", AMBER),
        ("No traceability", "Advice arrives by word of mouth with no record of the "
         "evidence used, who verified it, or whether the crop recovered. Nothing "
         "is auditable, so nothing improves.", GRAY),
    ]
    w = (SW - 2 * M - 0.3 * 2) / 3
    for i, (t, d, accent) in enumerate(cards):
        x = M + i * (w + 0.3)
        card(s, x, 2.62, w, 3.0)
        shape(s, x + 0.28, 2.88, 0.5, 0.075, fill=accent)
        tf = text(s, x + 0.28, 3.08, w - 0.56, 0.6)
        par(tf, t, size=15, color=SLATE, bold=True, font=HEAD, first=True, after=0)
        tf = text(s, x + 0.28, 3.62, w - 0.56, 1.8)
        par(tf, d, size=11, color=GRAY, first=True, after=0, line=1.12)

    card(s, M, 5.92, SW - 2 * M, 0.85, fill=MINT, line=None)
    tf = text(s, M + 0.35, 6.06, SW - 2 * M - 0.7, 0.6)
    par(tf, "The agronomic knowledge exists \u2014 ICAR / IRRI protocols are well "
            "documented. The failure is delivery: it does not reach the farmer at "
            "the moment of need, in a form that is actionable and safe.",
        size=12, color=RGBColor(0x06, 0x5F, 0x46), bold=True, first=True, after=0,
        line=1.08)
    footer(s)
    notes(s, "Frame the problem as a delivery and trust failure, not a knowledge "
             "gap. Land the idea that unsafe advice is worse than no advice.")
    return s


def slide_invariant(prs):
    s = slide(prs)
    bg(s, WHITE)
    header(s, "Design Principle", "The Non-Negotiable Invariant")

    card(s, M, 1.48, SW - 2 * M, 1.72, fill=DARK, line=None)
    shape(s, M, 1.48, 0.09, 1.72, fill=EMERALD)
    tf = text(s, M + 0.45, 1.70, SW - 2 * M - 0.9, 1.3, anchor=MSO_ANCHOR.MIDDLE)
    par(tf, "\u201cData and validated agricultural rules determine risk and allowed "
            "actions. AI explains validated results; AI does not independently "
            "decide what a farmer should do.\u201d",
        size=17, color=WHITE, bold=True, font=HEAD, first=True, after=0, line=1.12)

    left = M
    right = M + (SW - 2 * M) / 2 + 0.15
    cw = (SW - 2 * M) / 2 - 0.15

    card(s, left, 3.45, cw, 2.55, fill=MINT, line=None)
    tf = text(s, left + 0.32, 3.66, cw - 0.64, 0.5)
    par(tf, "\u2705  The deterministic core", size=15, color=RGBColor(0x06, 0x5F, 0x46),
        bold=True, font=HEAD, first=True, after=0)
    tf = text(s, left + 0.32, 4.20, cw - 0.64, 1.7)
    bullets(tf, [
        "Versioned rule set evaluates evidence \u2192 severity.",
        "Pre-stored action library keyed by ruleFamily::riskLevel.",
        "Safety gate validates every plan before display.",
        "Same inputs \u2192 identical plan, every single time.",
    ], size=11.5, color=RGBColor(0x13, 0x3B, 0x27), first=True, after=7)

    card(s, right, 3.45, cw, 2.55, fill=LIGHT, line=None)
    tf = text(s, right + 0.32, 3.66, cw - 0.64, 0.5)
    par(tf, "\U0001F916  The AI layer (subordinate)", size=15, color=SLATE,
        bold=True, font=HEAD, first=True, after=0)
    tf = text(s, right + 0.32, 4.20, cw - 0.64, 1.7)
    bullets(tf, [
        "Gemini runs strictly after risk + safety decisions.",
        "Strict prompt boundary: explain only, never classify.",
        "Cannot alter severity, actions or safety interceptions.",
        "Deterministic fallback \u2014 the app works with Gemini offline.",
    ], size=11.5, color=SLATE, first=True, after=7)

    card(s, M, 6.18, SW - 2 * M, 0.62, fill=WHITE, line=BORDER)
    tf = text(s, M + 0.3, 6.30, SW - 2 * M - 0.6, 0.4)
    par(tf, "Consequence:  the platform stays useful, auditable and defensible "
            "even when the model is unavailable, wrong or over-eager.",
        size=11.5, color=SLATE, italic=True, first=True, after=0)
    footer(s)
    notes(s, "This is the credibility slide. Explain WHY: agricultural advice is "
             "safety-critical, so a stochastic decision-maker is unacceptable. "
             "Determinism also makes the system testable.")
    return s


def slide_architecture(prs):
    s = slide(prs)
    bg(s, WHITE)
    header(s, "System Design", "Architecture & Decision Pipeline")

    stages = [
        ("Inputs", "Live weather (Open-Meteo + calibrated agro-climatic fallback), "
                   "farmer observations, ground measurements."),
        ("Evidence Engine", "Normalizes every signal into typed evidence carrying "
                            "source, strength, reliability and freshness."),
        ("Deterministic Risk Engine", "Versioned rule set maps evidence \u2192 severity "
                                       "(NORMAL \u2192 CRITICAL) with no randomness."),
        ("Safety Gate", "Validates the planned actions and injects AVOID "
                        "interceptions for hazardous interventions."),
        ("Action Plan Library", "Returns a pre-stored expert plan: DO NOW, DO NEXT, "
                                "AVOID, MONITOR, ESCALATE."),
        ("Explain & Verify", "AI explains the validated result; field workers verify "
                             "ground truth; outcome is logged."),
    ]
    y0 = 1.50
    rh = 0.72
    for i, (t, d) in enumerate(stages):
        y = y0 + i * rh
        if i < len(stages) - 1:
            shape(s, 0.905, y + 0.46, 0.02, rh - 0.06, fill=BORDER)
        shape(s, 0.72, y + 0.09, 0.42, 0.42, fill=GREEN,
              kind=MSO_SHAPE.OVAL)
        tfn = text(s, 0.72, y + 0.155, 0.42, 0.3)
        par(tfn, str(i + 1), size=13, color=WHITE, bold=True, font=HEAD,
            align=PP_ALIGN.CENTER, first=True, after=0)
        tf = text(s, 1.40, y + 0.05, SW - 2 * M - 0.9, 0.62)
        par(tf, t, size=13.5, color=SLATE, bold=True, font=HEAD, first=True,
            after=1)
        par(tf, d, size=10.5, color=GRAY, after=0, line=1.0)

    card(s, M, 5.90, SW - 2 * M, 0.88, fill=DARK, line=None)
    tf = text(s, M + 0.34, 6.06, SW - 2 * M - 0.68, 0.6)
    par(tf, "Persistence layer  \u00b7  MongoDB hydrate-on-boot, write-behind "
            "(debounced \u2248800 ms)  \u00b7  transparent in-memory fallback if the "
            "database is unreachable  \u00b7  live health at GET /api/system/persistence",
        size=11.5, color=PALE, first=True, after=0, line=1.05)
    footer(s)
    notes(s, "Walk the pipeline top to bottom in one breath. Stress that the AI "
             "step sits near the END, after every safety decision is already made.")
    return s


def slide_roles(prs):
    s = slide(prs)
    bg(s, WHITE)
    header(s, "Users", "Roles & Personas")

    roles = [
        ("FARMER", "Report incidents; view farm status, live weather and active "
                   "incidents; follow the validated plan; add ground measurements; "
                   "run the scenario simulator; submit outcome feedback.",
         ["/", "/incidents/new", "/incidents/:id", "/history"], GREEN),
        ("FIELD_WORKER", "See a queue of incidents flagged for verification, with "
                         "distance indicators; submit ground-truth verification and "
                         "optionally override severity.",
         ["/field-worker"], EMERALD),
        ("FPO_ADMIN", "Planned role for FPO-level coordination \u2014 community "
                      "drainage canals, pump support, seed and nursery logistics. "
                      "Today it appears only in escalation text.",
         ["\u2014 (planned)"], AMBER),
        ("SYSTEM_ADMIN", "Inspect persistence and platform health via the system "
                         "API.",
         ["GET /api/system/persistence"], GRAY),
    ]
    w = (SW - 2 * M - 0.28) / 2
    for i, (name, desc, screens, accent) in enumerate(roles):
        x = M + (i % 2) * (w + 0.28)
        y = 1.52 + (i // 2) * 2.35
        card(s, x, y, w, 2.1)
        shape(s, x, y, w, 0.075, fill=accent)
        tf = text(s, x + 0.3, y + 0.26, w - 0.6, 0.4)
        par(tf, name, size=14.5, color=SLATE, bold=True, font=HEAD, first=True,
            after=0)
        tf = text(s, x + 0.3, y + 0.70, w - 0.6, 1.0)
        par(tf, desc, size=11, color=GRAY, first=True, after=0, line=1.08)
        tf = text(s, x + 0.3, y + 1.68, w - 0.6, 0.3)
        par(tf, "Screens:  " + "  ".join(screens), size=9.5, color=accent,
            bold=True, first=True, after=0)

    card(s, M, 6.22, SW - 2 * M, 0.58, fill=LIGHT, line=None)
    tf = text(s, M + 0.3, 6.33, SW - 2 * M - 0.6, 0.4)
    par(tf, "Demo personas:  Krishna Reddy \u2014 farmer, 2.8-acre paddy "
            "(\u201cVemuru Riverbed Field\u201d, BPT 5204 / Samba Mahsuri, tillering, "
            "poor drainage).    Ramesh Kumar \u2014 agricultural field officer.",
        size=10.5, color=SLATE, first=True, after=0)
    footer(s)
    notes(s, "Be explicit that authentication is NOT implemented \u2014 roles exist in "
             "the type system and UI. Say it before someone finds it.")
    return s


def slide_features(prs):
    s = slide(prs)
    bg(s, WHITE)
    header(s, "Product", "Core Capabilities")

    feats = [
        ("Farmer Dashboard", "Live 24h rainfall, temperature, humidity and wind "
                             "with active incidents and severity badges."),
        ("Emergency Reporting Wizard", "Non-academic 3-step flow: condition, "
                                       "symptom chips, evidence confirmation."),
        ("Hero Action Screen", "DO NOW / AVOID / MONITOR, why-this-recommendation, "
                               "evidence traceability, live reassessment."),
        ("Field Worker Queue", "Human-in-the-loop verification: confirm, partially "
                               "confirm, reject or request more information."),
        ("Auditable History", "Non-destructive timeline preserving every "
                              "assessment version, action plan and outcome."),
        ("Scenario Simulator", "Manipulate rainfall, duration, depth, drainage and "
                               "crop stage \u2192 instant re-assessment everywhere."),
    ]
    w = (SW - 2 * M - 0.28 * 2) / 3
    h = 1.72
    for i, (t, d) in enumerate(feats):
        x = M + (i % 3) * (w + 0.28)
        y = 1.52 + (i // 3) * (h + 0.26)
        card(s, x, y, w, h)
        shape(s, x + 0.26, y + 0.26, 0.42, 0.42, fill=MINT,
              kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.25)
        tfn = text(s, x + 0.26, y + 0.335, 0.42, 0.3)
        par(tfn, str(i + 1), size=12, color=GREEN, bold=True, font=HEAD,
            align=PP_ALIGN.CENTER, first=True, after=0)
        tf = text(s, x + 0.26, y + 0.78, w - 0.52, 0.4)
        par(tf, t, size=12.5, color=SLATE, bold=True, font=HEAD, first=True,
            after=0)
        tf = text(s, x + 0.26, y + 1.14, w - 0.52, 0.55)
        par(tf, d, size=10, color=GRAY, first=True, after=0, line=1.05)

    tf = text(s, M, 5.62, SW - 2 * M, 0.35)
    par(tf, "Plus, on every screen:", size=11.5, color=SLATE, bold=True,
        font=HEAD, first=True, after=0)
    tf = text(s, M, 5.94, SW - 2 * M, 0.9)
    bullets(tf, [
        "Pre-stored recommendation library \u2014 12 expert plans, delivered instantly "
        "(no generation step, identical on every recurrence).",
        "Optional durable persistence \u2014 set MONGODB_URI and restarts survive; "
        "otherwise the app runs transparently in memory.",
    ], size=11, color=GRAY, first=True, after=5)
    footer(s)
    notes(s, "Do not read every card. Pick two or three and spend the time on the "
             "live demo instead.")
    return s


def slide_hero_screen(prs):
    s = slide(prs)
    bg(s, WHITE)
    header(s, "Core UI", "The Emergency Action Screen")

    tf = text(s, M, 1.45, 5.5, 5.0)
    bullets(tf, [
        ("DO NOW", True),
        "Step-by-step immediate physical interventions (e.g. cut the drainage "
        "breach at the lowest bund corner).",
        ("AVOID (safety interceptions)", True),
        "Harmful actions the Safety Gate blocks \u2014 e.g. no granular urea into "
        "standing flood water, no contact pesticide sprays.",
        ("MONITOR", True),
        "Symptoms and submergence thresholds to watch, with a review deadline.",
        ("WHY THIS RECOMMENDATION?", True),
        "Plain-language explanation generated strictly from the validated result.",
        ("Evidence view", True),
        "Traceability cards showing source, reliability and freshness "
        "(\U0001F7E2 Fresh  \U0001F7E1 Aging  \U0001F534 Stale).",
        ("Reassess or resolve", True),
        "Add new ground measurements for live recalculation, then submit the "
        "outcome.",
    ], size=11, color=SLATE, first=True, after=6)

    img_fit(s, os.path.join(SHOTS, "actions.png"), 6.35, 1.75, 6.35, 4.15)
    tf = text(s, 6.35, 6.05, 6.35, 0.5)
    par(tf, "Incident response screen \u2014 /incidents/:id (live build)",
        size=9.5, color=GRAY, italic=True, align=PP_ALIGN.CENTER, first=True,
        after=0)
    footer(s)
    notes(s, "This is the money screen. If demoing live, show DO NOW, then AVOID, "
             "then open the Simulator drawer.")
    return s


def slide_dashboard_wizard(prs):
    s = slide(prs)
    bg(s, WHITE)
    header(s, "Product Tour", "Dashboard & Emergency Reporting")

    img_fit(s, os.path.join(SHOTS, "dashboard.png"), M, 1.72, 5.9, 3.55)
    tf = text(s, M, 5.38, 5.9, 1.1)
    par(tf, "Farmer Dashboard  \u2014  /", size=12, color=SLATE, bold=True,
        font=HEAD, align=PP_ALIGN.CENTER, first=True, after=3)
    par(tf, "Live environmental conditions, active incidents with severity badges, "
            "and a prominent emergency CTA: \u201cSomething wrong with your farm right "
            "now?\u201d", size=10, color=GRAY, align=PP_ALIGN.CENTER, after=0,
        line=1.05)

    x2 = M + 5.9 + 0.4
    w2 = SW - M - x2
    img_fit(s, os.path.join(SHOTS, "wizard.png"), x2, 1.72, w2, 3.55)
    tf = text(s, x2, 5.38, w2, 1.1)
    par(tf, "Reporting Wizard  \u2014  /incidents/new", size=12, color=SLATE,
        bold=True, font=HEAD, align=PP_ALIGN.CENTER, first=True, after=3)
    par(tf, "Three steps: what happened, quick symptom chips (e.g. \u201cwater "
            "standing > 15 cm\u201d), then \u201cLet\u2019s check what we know\u201d before the risk "
            "engine runs.", size=10, color=GRAY, align=PP_ALIGN.CENTER, after=0,
        line=1.05)
    footer(s)
    notes(s, "The wizard is deliberately non-academic \u2014 the farmer never needs the "
             "scientific name of the problem.")
    return s


def slide_risk_engine(prs):
    s = slide(prs)
    bg(s, WHITE)
    header(s, "Intelligence", "Deterministic Risk Engine")

    tf = text(s, M, 1.42, SW - 2 * M, 0.55)
    par(tf, "Paddy waterlogging thresholds (as implemented)",
        size=13, color=SLATE, bold=True, font=HEAD, first=True, after=0)

    rows = [
        ("CRITICAL", "rainfall \u2265 100 mm  OR  (rainfall \u2265 75 mm AND standing "
                     "water \u2265 36 h)  \u2014 both require poor drainage", RED, True),
        ("HIGH", "rainfall \u2265 60 mm  OR  standing water \u2265 24 h  OR  depth \u2265 15 cm  "
                 "OR  rainfall \u2265 50 mm in a sensitive stage", RGBColor(0xC2, 0x41, 0x0C), True),
        ("MODERATE", "rainfall \u2265 30 mm  OR  standing water \u2265 8 h", AMBER, False),
        ("LOW", "rainfall \u2265 15 mm  OR  any standing water > 0 h", RGBColor(0x94, 0xA3, 0xB8), False),
        ("NORMAL", "otherwise", RGBColor(0x64, 0x74, 0x8B), False),
    ]
    y = 1.92
    for name, cond, accent, esc in rows:
        card(s, M, y, SW - 2 * M, 0.66, fill=LIGHT, line=None, radius=0.12)
        shape(s, M + 0.14, y + 0.14, 1.15, 0.38, fill=accent,
              kind=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.28)
        tfn = text(s, M + 0.14, y + 0.205, 1.15, 0.3)
        par(tfn, name, size=10, color=WHITE, bold=True, font=HEAD,
            align=PP_ALIGN.CENTER, first=True, after=0)
        tf = text(s, M + 1.48, y + 0.20, SW - 2 * M - 3.2, 0.32)
        par(tf, cond, size=10.5, color=SLATE, first=True, after=0)
        tf = text(s, SW - M - 1.6, y + 0.20, 1.4, 0.32)
        par(tf, "Escalates" if esc else "No escalation", size=9.5,
            color=accent if esc else GRAY, bold=esc, align=PP_ALIGN.RIGHT,
            first=True, after=0)
        y += 0.74

    card(s, M, 5.72, SW - 2 * M, 1.0, fill=MINT, line=None)
    tf = text(s, M + 0.32, 5.84, SW - 2 * M - 0.64, 0.78)
    par(tf, "Rule-priority policy:  an explicitly reported incident type always "
            "routes to its dedicated rule. Ambient signals (heavy rain, standing "
            "water, extreme heat) only auto-route \u2018Something else\u2019 reports \u2014 so a "
            "heat-stress report can never be hijacked into a waterlogging plan.",
        size=10.5, color=RGBColor(0x06, 0x5F, 0x46), first=True, after=4, line=1.06)
    par(tf, "Review cadence:  reviewAt = assessedAt + { CRITICAL: 3 h, HIGH: 6 h, "
            "otherwise: 12 h }.", size=10.5, color=RGBColor(0x06, 0x5F, 0x46),
        bold=True, after=0)
    footer(s)
    notes(s, "Heat stress adds CRITICAL at \u226540\u00b0C during flowering / panicle "
             "initiation. Pest/disease deliberately never recommends a pesticide.")
    return s


def slide_library(prs):
    s = slide(prs)
    bg(s, WHITE)
    header(s, "Content", "Recommendation Library & Safety Gate")

    card(s, M, 1.48, 6.05, 2.35, fill=LIGHT, line=None)
    tf = text(s, M + 0.3, 1.66, 5.45, 0.4)
    par(tf, "12 pre-stored expert action plans", size=13.5, color=SLATE,
        bold=True, font=HEAD, first=True, after=0)
    tf = text(s, M + 0.3, 2.10, 5.45, 1.6)
    bullets(tf, [
        "Waterlogging \u2014 5 severities (NORMAL \u2192 CRITICAL).",
        "Heat stress and pest / disease cover every severity the engine emits.",
        "Keyed by ruleFamily::riskLevel \u2014 plans appear instantly, with no "
        "generation step.",
        "Identical plan every time the same conditions recur (deterministic).",
    ], size=11, color=SLATE, first=True, after=6)

    card(s, 6.95, 1.48, SW - M - 6.95, 2.35, fill=DARK, line=None)
    tf = text(s, 7.25, 1.66, SW - M - 7.55, 0.4)
    par(tf, "Safety Gate", size=13.5, color=WHITE, bold=True, font=HEAD,
        first=True, after=0)
    tf = text(s, 7.25, 2.10, SW - M - 7.55, 1.6)
    bullets(tf, [
        "Validates every plan before it is displayed.",
        "Blocks hazardous interventions outright (AVOID).",
        "Escalates to a human when evidence is thin or conflicting.",
        "The engine never recommends a pesticide \u2014 verified diagnosis first.",
    ], size=11, color=PALE, first=True, after=6)

    tf = text(s, M, 3.98, SW - 2 * M, 0.35)
    par(tf, "Representative interceptions", size=12, color=SLATE, bold=True,
        font=HEAD, first=True, after=0)
    examples = [
        "DO NOT apply granular urea into standing flood water \u2014 it is lost and "
        "worsens the anaerobic condition.",
        "DO NOT spray contact pesticides on a water-stressed or submerged canopy \u2014 "
        "phytotoxicity risk.",
        "DO NOT re-transplant or top-dress until standing water has receded and "
        "the review deadline has passed.",
    ]
    w = (SW - 2 * M - 0.3 * 2) / 3
    for i, e in enumerate(examples):
        x = M + i * (w + 0.3)
        card(s, x, 4.36, w, 1.72, fill=WHITE, line=BORDER)
        shape(s, x, 4.36, w, 0.06, fill=RED)
        tf = text(s, x + 0.24, 4.56, w - 0.48, 1.4)
        par(tf, e, size=10.5, color=SLATE, first=True, after=0, line=1.1)
    footer(s)
    notes(s, "The safety gate is what makes this credible to an agronomist. It is "
             "also the clearest demonstration that AI does not decide.")
    return s


def slide_evidence(prs):
    s = slide(prs)
    bg(s, WHITE)
    header(s, "Trust", "Evidence Quality & Human Verification")

    tf = text(s, M, 1.45, 5.9, 4.0)
    par(tf, "Every assessment is traceable", size=13.5, color=SLATE, bold=True,
        font=HEAD, first=True, after=6)
    bullets(tf, [
        "Each observation carries source, reliability, strength and timestamps "
        "(observedAt / retrievedAt / freshnessMinutes).",
        "Freshness is surfaced, never hidden: \U0001F7E2 Fresh  \U0001F7E1 Aging  "
        "\U0001F534 Stale.",
        "Observed vs forecast values are labelled distinctly \u2014 no silent blending.",
        "Unified evidence view returns verdicts plus what is still missing.",
    ], size=11, color=SLATE, first=False, after=7)

    tf = text(s, M, 4.05, 5.9, 2.4)
    par(tf, "Human-in-the-loop escalation", size=13.5, color=SLATE, bold=True,
        font=HEAD, first=True, after=6)
    bullets(tf, [
        "HIGH / CRITICAL incidents auto-escalate to the field-worker queue.",
        "Officer confirms, partially confirms, rejects or requests more "
        "information, and may override severity.",
        "Verified ground truth is appended to the audit trail, not overwritten.",
    ], size=11, color=SLATE, first=False, after=7)

    img_fit(s, os.path.join(SHOTS, "fieldworker.png"), 6.75, 1.62, SW - M - 6.75, 3.7)
    tf = text(s, 6.75, 5.42, SW - M - 6.75, 1.0)
    par(tf, "Field Worker & Extension Queue \u2014 /field-worker", size=11.5,
        color=SLATE, bold=True, font=HEAD, align=PP_ALIGN.CENTER, first=True,
        after=3)
    par(tf, "Pending inspections with distance indicators and a structured "
            "verification form.", size=10, color=GRAY, align=PP_ALIGN.CENTER,
        after=0, line=1.05)
    footer(s)
    notes(s, "Emphasise the design intent: the system knows what it does not know, "
             "and it hands uncertainty to a human rather than guessing.")
    return s


def slide_persistence(prs):
    s = slide(prs)
    bg(s, WHITE)
    header(s, "Reliability", "Persistence, Degradation & Audit")

    img_fit(s, os.path.join(SHOTS, "history.png"), M, 1.66, 6.0, 3.9)
    tf = text(s, M, 5.68, 6.0, 0.9)
    par(tf, "Auditable History \u2014 /history", size=11.5, color=SLATE,
        bold=True, font=HEAD, align=PP_ALIGN.CENTER, first=True, after=3)
    par(tf, "Non-destructive audit log preserving every assessment version, action "
            "plan, verification and farmer outcome.", size=10, color=GRAY,
        align=PP_ALIGN.CENTER, after=0, line=1.05)

    x2 = M + 6.35
    w2 = SW - M - x2
    tf = text(s, x2, 1.62, w2, 3.9)
    par(tf, "Durable by default, graceful by design", size=13, color=SLATE,
        bold=True, font=HEAD, first=True, after=8)
    bullets(tf, [
        "Set MONGODB_URI and every incident / farm mutation is persisted.",
        "Write-behind, debounced \u2248800 ms after each mutation.",
        "Boot hydrates from the database \u2014 restarts survive.",
        "If the database is unreachable the app transparently runs in memory and "
        "retries in the background.",
        "Live health: GET /api/system/persistence.",
        "Reassessments append; nothing is destroyed or silently rewritten.",
    ], size=11, color=SLATE, first=False, after=9)
    footer(s)
    notes(s, "Honesty note to be ready for: two dev servers running against the "
             "same database overwrote each other's data. One writer at a time.")
    return s


def slide_tech(prs):
    s = slide(prs)
    bg(s, WHITE)
    header(s, "Engineering", "Technology Stack")

    groups = [
        ("Application", ["Next.js 14 \u2014 App Router", "TypeScript (strict)",
                         "Tailwind CSS", "Lucide React icons"]),
        ("Data & Decisions", ["Versioned deterministic rule engine",
                              "12-plan expert action library",
                              "Evidence freshness & verdict pipeline",
                              "Safety gate interceptions"]),
        ("External Services", ["Open-Meteo live weather API",
                               "Calibrated agro-climatic fallback model",
                               "Google Gemini \u2014 explanation layer only",
                               "MongoDB \u2014 optional durable store"]),
        ("Quality", ["Zero-dependency Node test runner",
                     "31 tests passing",
                     "Strict prompt boundary on the AI layer",
                     "ICAR / IRRI-aligned agronomic content"]),
    ]
    w = (SW - 2 * M - 0.3) / 2
    for i, (title, items) in enumerate(groups):
        x = M + (i % 2) * (w + 0.3)
        y = 1.55 + (i // 2) * 2.42
        card(s, x, y, w, 2.18)
        shape(s, x, y, 0.07, 2.18, fill=GREEN)
        tf = text(s, x + 0.32, y + 0.24, w - 0.64, 0.4)
        par(tf, title, size=13, color=SLATE, bold=True, font=HEAD, first=True,
            after=0)
        tf = text(s, x + 0.32, y + 0.68, w - 0.64, 1.4)
        for j, it in enumerate(items):
            par(tf, "\u2022  " + it, size=10.5, color=GRAY, after=6, first=(j == 0),
                line=1.02)

    card(s, M, 6.42, SW - 2 * M, 0.5, fill=MINT, line=None)
    tf = text(s, M + 0.3, 6.53, SW - 2 * M - 0.6, 0.32)
    par(tf, "No exotic dependencies, no paid services required to run the demo \u2014 "
            "GEMINI_API_KEY and MONGODB_URI are both optional.",
        size=10.5, color=RGBColor(0x06, 0x5F, 0x46), bold=True, first=True, after=0)
    footer(s)
    notes(s, "If asked about scale: the persistence layer was the last milestone "
             "and is deliberately swappable behind a store interface.")
    return s


def slide_demo(prs):
    s = slide(prs)
    bg(s, WHITE)
    header(s, "Walkthrough", "Primary Demonstration Script")

    steps = [
        ("Baseline", "Krishna Reddy\u2019s paddy field in tillering faces 82 mm of rain "
                     "and 24 h of standing water \u2192 HIGH RISK."),
        ("Review the plan", "Inspect DO NOW (breach bunds, clear drainage), AVOID "
                            "(urea application blocked) and Why This Recommendation."),
        ("Field verification", "Open the Field Worker Queue, inspect the incident and "
                               "submit official verification."),
        ("Simulate escalation", "In the Scenario Simulator push rainfall to 115 mm "
                                "and duration to 48 h \u2192 instant CRITICAL RISK."),
        ("Simulate de-escalation", "Drop rainfall to 5 mm, standing water to 0 h and "
                                   "set drainage to good \u2192 LOW / NORMAL."),
        ("Report and persist", "Report a new incident through the wizard and watch "
                               "the deterministic engine assess it live; restart the "
                               "server to prove persistence."),
    ]
    y = 1.50
    for i, (t, d) in enumerate(steps):
        card(s, M, y, SW - 2 * M, 0.78, fill=LIGHT if i % 2 == 0 else WHITE,
             line=None if i % 2 == 0 else BORDER)
        shape(s, M + 0.16, y + 0.17, 0.44, 0.44, fill=GREEN, kind=MSO_SHAPE.OVAL)
        tfn = text(s, M + 0.16, y + 0.235, 0.44, 0.3)
        par(tfn, str(i + 1), size=12.5, color=WHITE, bold=True, font=HEAD,
            align=PP_ALIGN.CENTER, first=True, after=0)
        tf = text(s, M + 0.82, y + 0.10, SW - 2 * M - 1.1, 0.62)
        par(tf, t, size=12, color=SLATE, bold=True, font=HEAD, first=True, after=1)
        par(tf, d, size=10, color=GRAY, after=0, line=1.0)
        y += 0.845

    tf = text(s, M, 6.62, SW - 2 * M, 0.35)
    par(tf, "Run the demo from a single dev server. Only one process should own the "
            "shared Atlas data at a time.", size=10, color=GRAY, italic=True,
        first=True, after=0)
    footer(s)
    notes(s, "Keep the demo under four minutes. The escalation / de-escalation pair "
             "is the strongest thirty seconds \u2014 it proves determinism live.")
    return s


def slide_parked(prs):
    s = slide(prs)
    bg(s, WHITE)
    header(s, "Honest Scope", "Parked for the Demo Build")

    card(s, M, 1.48, SW - 2 * M, 1.0, fill=RGBColor(0xFE, 0xF3, 0xC7), line=None)
    tf = text(s, M + 0.34, 1.62, SW - 2 * M - 0.68, 0.75)
    par(tf, "\u26A0  Field Map (/field) and Plot Registry (/plots) are disabled in the "
            "demo build \u2014 they return 404 and are not linked in the navigation.",
        size=12.5, color=RGBColor(0x92, 0x40, 0x0E), bold=True, first=True, after=0,
        line=1.06)

    tf = text(s, M, 2.68, 6.2, 3.4)
    par(tf, "Why they were parked", size=13, color=SLATE, bold=True, font=HEAD,
        first=True, after=6)
    bullets(tf, [
        "The registry could only find plots that had been registered inside "
        "FarmRescue \u2014 there is no government land-records verification, so it "
        "made a promise the data could not keep.",
        "Plot coordinates came from user-placed map taps or GPS, not cadastral "
        "survey data.",
        "Two concurrent dev servers overwrote each other in the shared database, "
        "making the feature look broken in front of evaluators.",
    ], size=11, color=GRAY, first=False, after=8)

    tf = text(s, 7.1, 2.68, SW - M - 7.1, 3.4)
    par(tf, "Status and path forward", size=13, color=SLATE, bold=True,
        font=HEAD, first=True, after=6)
    bullets(tf, [
        "The full implementation is preserved, not deleted \u2014 cascading "
        "State \u2192 Village \u2192 Khasra lookup, map chips, register-on-the-spot, "
        "Bhuiyan-style khasra report entry.",
        "Restore is a one-step move from the parked folder.",
        "Revisit requires a real land-records API \u2014 a Bhulekh / ULPIN-style "
        "integration, plus a single-writer data policy.",
    ], size=11, color=GRAY, first=False, after=8)

    card(s, M, 6.15, SW - 2 * M, 0.62, fill=LIGHT, line=None)
    tf = text(s, M + 0.3, 6.27, SW - 2 * M - 0.6, 0.4)
    par(tf, "Decision: show a smaller honest product rather than a larger one that "
            "cannot be defended in a live demo.",
        size=11, color=SLATE, italic=True, first=True, after=0)
    footer(s)
    notes(s, "Say this unprompted. Volunteering the parked scope buys credibility "
             "for everything else in the deck.")
    return s


def slide_gaps_roadmap(prs):
    s = slide(prs)
    bg(s, WHITE)
    header(s, "Forward Look", "Known Gaps & Roadmap")

    tf = text(s, M, 1.42, 6.1, 0.4)
    par(tf, "Known gaps (honest list)", size=13, color=SLATE, bold=True,
        font=HEAD, first=True, after=0)
    tf = text(s, M, 1.84, 6.1, 4.3)
    bullets(tf, [
        "No authentication \u2014 roles exist in the type system and UI, but API "
        "routes accept no credentials.",
        "Recommendation library is paddy-centric for waterlogging; other crops "
        "fall through to the generic plan.",
        "Incident creation seeds a fixed 65 mm rainfall reading; standing water "
        "defaults to 12 h / 10 cm.",
        "\u201cNext reassessment scheduled\u201d is display-only \u2014 no background scheduler "
        "exists.",
        "Test coverage is partial: risk boundaries, safety-gate rules and library "
        "coverage are still untested.",
        "No real SMS / WhatsApp / IVR delivery \u2014 notification exists only as "
        "escalation text.",
    ], size=10.5, color=GRAY, first=False, after=7)

    x2 = M + 6.45
    w2 = SW - M - x2
    tf = text(s, x2, 1.42, w2, 0.4)
    par(tf, "Roadmap", size=13, color=SLATE, bold=True, font=HEAD, first=True,
        after=0)

    phases = [
        ("P1 \u00b7 Hardening", "Boundary tests for risk + safety rules; a real "
                               "scheduled reassessment loop; live weather evidence on "
                               "incident creation; basic role authentication.", GREEN),
        ("P2 \u00b7 Breadth", "Multi-crop action plans; photo evidence capture \u2192 "
                             "storage \u2192 verification; more agro-climatic regions.", EMERALD),
        ("P3 \u00b7 Depth", "Offline-first PWA; sensor / IoT intake; outcome learning "
                           "loop; multi-tenancy for FPOs; land-records integration.",
         RGBColor(0x0E, 0x74, 0x90)),
    ]
    y = 1.84
    for name, desc, accent in phases:
        card(s, x2, y, w2, 1.42)
        shape(s, x2, y, 0.07, 1.42, fill=accent)
        tf = text(s, x2 + 0.3, y + 0.16, w2 - 0.6, 0.35)
        par(tf, name, size=12, color=SLATE, bold=True, font=HEAD, first=True,
            after=0)
        tf = text(s, x2 + 0.3, y + 0.53, w2 - 0.6, 0.85)
        par(tf, desc, size=10, color=GRAY, first=True, after=0, line=1.08)
        y += 1.58

    card(s, M, 6.28, SW - 2 * M, 0.5, fill=MINT, line=None)
    tf = text(s, M + 0.3, 6.39, SW - 2 * M - 0.6, 0.32)
    par(tf, "The P0 MVP is complete and demo-ready. Everything above is scheduled "
            "work, not a hidden failure.", size=10.5,
        color=RGBColor(0x06, 0x5F, 0x46), bold=True, first=True, after=0)
    footer(s)
    notes(s, "Cite gaps from the PRD, not from memory \u2014 they are documented in "
             "section 14 and roadmap phases in section 15.")
    return s


def slide_close(prs):
    s = slide(prs)
    bg_gradient(s, FOREST, RGBColor(0x05, 0x1A, 0x10), 45)
    shape(s, 9.8, -1.6, 5.0, 5.0, fill=RGBColor(0x12, 0x4A, 0x2E),
          kind=MSO_SHAPE.OVAL)

    tf = text(s, M + 0.2, 1.15, 11.6, 0.5)
    par(tf, "CLOSING", size=12, color=EMERALD, bold=True, first=True, after=0)

    tf = text(s, M + 0.2, 1.70, 11.6, 2.0)
    par(tf, "When the farm goes wrong at 3 a.m.,", size=30, color=WHITE,
        bold=True, font=HEAD, first=True, after=2)
    par(tf, "FarmRescue already knows what to do.", size=30, color=EMERALD,
        bold=True, font=HEAD, after=0)

    shape(s, M + 0.2, 3.92, 1.0, 0.075, fill=EMERALD)

    tf = text(s, M + 0.2, 4.22, 11.6, 1.2)
    par(tf, "A deterministic agronomic decision engine, an expert-authored response "
            "library, a safety gate that blocks harmful advice, and a human "
            "verification loop \u2014 with AI strictly in the role of explainer.",
        size=13.5, color=PALE, first=True, after=0, line=1.12)

    tf = text(s, M + 0.2, 5.62, 11.6, 0.9)
    par(tf, "Thank you  \u00b7  Questions welcome", size=20, color=WHITE,
        bold=True, font=HEAD, first=True, after=4)
    par(tf, "FarmRescue \u2014 Emergency Agricultural Early-Warning & Response Platform "
            "\u00b7 October 2026", size=11, color=RGBColor(0x9C, 0xD8, 0xB8),
        after=0)
    notes(s, "Close on the determinism message, then invite questions \u2014 the "
             "questions will be about safety and trust, which is exactly the ground "
             "the system is built to defend.")
    return s


# ----------------------------------------------------------------------------
def main():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    prs.core_properties.title = "FarmRescue \u2014 Emergency Agricultural Early-Warning & Response Platform"
    prs.core_properties.author = "FarmRescue Project"
    prs.core_properties.subject = "Official project briefing"
    prs.core_properties.comments = ("Deterministic agronomic risk intelligence. "
                                    "Data and rules decide; AI explains.")

    slide_title(prs)
    slide_summary(prs)
    slide_problem(prs)
    slide_invariant(prs)
    slide_architecture(prs)
    slide_roles(prs)
    slide_features(prs)
    slide_hero_screen(prs)
    slide_dashboard_wizard(prs)
    slide_risk_engine(prs)
    slide_library(prs)
    slide_evidence(prs)
    slide_persistence(prs)
    slide_tech(prs)
    slide_demo(prs)
    slide_parked(prs)
    slide_gaps_roadmap(prs)
    slide_close(prs)

    prs.save(OUT)
    print("Wrote %s" % OUT)
    print("Slides: %d" % len(prs.slides._sldIdLst))


if __name__ == "__main__":
    main()
